import io

from app.config import get_settings
from app.main import app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

client = TestClient(app)


def test_templates_and_preview_and_generate():
    assert {t["id"] for t in client.get("/api/templates").json()} == {"default-id", "default-en"}
    r = client.post("/api/timesheet/preview", json={"year": 2026, "month": 1})
    assert r.status_code == 200
    draft = r.json()
    assert len(draft["rows"]) == 31
    g = client.post("/api/timesheet/generate", json={"draft": draft})
    assert g.status_code == 200
    assert "TimeSheet_Aria_Nur_Jamal_2026.xlsx" in g.headers["content-disposition"]
    assert load_workbook(io.BytesIO(g.content)).sheetnames == ["Januari"]


def test_invalid_manual_percent_returns_422():
    r = client.post(
        "/api/timesheet/preview",
        json={
            "year": 2026,
            "month": 1,
            "manual_percent": {
                "BPAS-5437": {"2026-01-21": 80, "2026-01-22": 40, "2026-01-23": 60, "2026-01-26": 100}
            },
        },
    )
    assert r.status_code == 422


def test_unknown_template_404():
    assert (
        client.post("/api/timesheet/preview", json={"year": 2026, "month": 1, "template_id": "x"}).status_code
        == 404
    )


def test_signature_upload_validation(tmp_path):
    get_settings().storage_dir = tmp_path
    buf = io.BytesIO()
    Image.new("RGB", (50, 20)).save(buf, "PNG")
    ok = client.post("/api/signatures/employee", files={"file": ("s.png", buf.getvalue(), "image/png")})
    assert ok.status_code == 204 and (tmp_path / "signatures" / "employee.png").exists()
    bad = client.post("/api/signatures/employee", files={"file": ("s.gif", b"GIF89a", "image/gif")})
    assert bad.status_code == 415
    big = client.post("/api/signatures/team_lead", files={"file": ("s.png", b"0" * 1_000_001, "image/png")})
    assert big.status_code == 413


def test_signature_status_and_delete(tmp_path):
    get_settings().storage_dir = tmp_path
    status_before = client.get("/api/signatures").json()
    assert status_before["employee"] is False

    buf = io.BytesIO()
    Image.new("RGB", (50, 20)).save(buf, "PNG")
    client.post("/api/signatures/employee", files={"file": ("s.png", buf.getvalue(), "image/png")})

    status_after = client.get("/api/signatures").json()
    assert status_after["employee"] is True

    del_res = client.delete("/api/signatures/employee")
    assert del_res.status_code == 204
    assert client.get("/api/signatures").json()["employee"] is False


def test_jira_status_and_formula_simulate():
    status = client.get("/api/jira/status").json()
    assert status["connected"] is True
    assert "mock" in status["site_url"] or "atlassian" in status["site_url"]

    sim = client.post(
        "/api/formula/simulate",
        json={
            "estimate_mandays": 4.0,
            "active_days": 3,
            "done": True,
            "start_percent": 0,
        },
    ).json()
    assert sim["total_progress"] == 100
    assert len(sim["daily_percentages"]) == 3
    assert sim["daily_percentages"][-1] == 100


def test_template_crud(tmp_path):
    new_tpl = {
        "id": "test-company",
        "name": "Test Company Template",
        "language": "id",
        "company_name": "PT. Uji Coba",
        "client": "Klien Uji",
        "department": "Engineering",
        "title": "Software Engineer",
        "employee_name": "Budi Santoso",
        "team_lead_name": "Siti Rahma",
        "work_start": "09:00",
        "work_end": "18:00",
        "break_minutes": 60,
        "show_all_days": True,
        "timezone": "Asia/Jakarta",
        "holiday_country": "ID",
        "extra_holidays": ["2026-05-01"],
        "file_name_pattern": "TimeSheet_{name}_{year}.xlsx",
        "formula": {
            "hours_per_manday": 8,
            "story_point_to_manday": 1,
            "max_progress_before_done": 95,
            "rounding_step": 5,
            "min_daily_portion": 5,
            "segmentation_mode": "random",
            "carry_over_between_months": True,
        },
    }
    try:
        create_res = client.post("/api/templates", json=new_tpl)
        assert create_res.status_code == 201
        assert client.get("/api/templates/test-company").status_code == 200

        new_tpl["company_name"] = "PT. Uji Coba Diperbarui"
        update_res = client.put("/api/templates/test-company", json=new_tpl)
        assert update_res.status_code == 200
        assert update_res.json()["company_name"] == "PT. Uji Coba Diperbarui"
    finally:
        del_res = client.delete("/api/templates/test-company")
        assert del_res.status_code == 204
        assert client.get("/api/templates/test-company").status_code == 404


def test_multi_user_session_and_isolation(tmp_path):
    from app.api.routes import SESSIONS, UserSession

    get_settings().storage_dir = tmp_path

    # Daftarkan session untuk User 1 (Budi) dan User 2 (Siti)
    token1 = "token-user-budi"
    token2 = "token-user-siti"
    SESSIONS[token1] = UserSession(
        token=token1,
        user_id="budi-id",
        username="budi",
        email="budi@company.com",
        api_token="token123",
        site_url="https://jira.company.com",
        display_name="Budi Santoso",
        role="developer",
        template_id="default-id",
        account_id="budi-id",
    )
    SESSIONS[token2] = UserSession(
        token=token2,
        user_id="siti-id",
        username="siti",
        email="siti@company.com",
        api_token="token456",
        site_url="https://jira.company.com",
        display_name="Siti Rahma",
        role="developer",
        template_id="default-id",
        account_id="siti-id",
    )

    # Cek status user 1
    status1 = client.get("/api/jira/status", headers={"Authorization": f"Bearer {token1}"}).json()
    assert status1["display_name"] == "Budi Santoso"
    assert status1["email"] == "budi@company.com"

    # Cek status user 2
    status2 = client.get("/api/jira/status", headers={"Authorization": f"Bearer {token2}"}).json()
    assert status2["display_name"] == "Siti Rahma"
    assert status2["email"] == "siti@company.com"

    # Upload tanda tangan User 1
    buf = io.BytesIO()
    Image.new("RGB", (50, 20)).save(buf, "PNG")
    client.post(
        "/api/signatures/employee",
        files={"file": ("s.png", buf.getvalue(), "image/png")},
        headers={"Authorization": f"Bearer {token1}"},
    )
    # User 1 punya tanda tangan, tapi User 2 belum
    assert (
        client.get("/api/signatures", headers={"Authorization": f"Bearer {token1}"}).json()["employee"]
        is True
    )
    assert (
        client.get("/api/signatures", headers={"Authorization": f"Bearer {token2}"}).json()["employee"]
        is False
    )

    # Logout user 1
    logout = client.post("/api/auth/logout", headers={"Authorization": f"Bearer {token1}"})
    assert logout.status_code == 204
    assert token1 not in SESSIONS
