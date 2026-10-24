"""Render timesheet ke .xlsx dengan layout identik file referensi (Bagian 5 PRD)."""

from __future__ import annotations

import io
import math
from datetime import date, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage

from timesheet_core import i18n
from timesheet_core.preview import default_working_hours
from timesheet_core.schemas import DayRow, TemplateSpec, TimesheetDraft

FONT = "Arial"
RED = "FFFF0000"
BLACK = "FF000000"
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

COL_WIDTHS = {"A": 6.86, "B": 13.71, "C": 12.0, "D": 14.43, "E": 14.0, "F": 95.0}
HEADER_ROW_INFO = 11
HEADER_ROW_TABLE = 14
FIRST_DATA_ROW = 15
SIGNATURE_MAX_PX = (160, 70)
ASSETS_DIR = Path(__file__).resolve().parent / "templates" / "assets"


def _t(value: str) -> time:
    h, m = value.split(":")
    return time(int(h), int(m))


def info_text(row: DayRow) -> str:
    """`[Label] KEY: Judul (xx%) - aktivitas`, tiap task ke-2 dst. dipisah baris baru (\n)."""
    blocks: list[str] = []
    for s in row.stories:
        head = f"{s.key}: {s.title}"
        if s.epic:
            head = f"[{s.epic}] {head}"
        if s.percent is not None:
            head += f" ({s.percent}%)"
        blocks.append(head + "".join(f" - {a}" for a in s.activities))
    if row.free_text:
        blocks.append(row.free_text)
    return "\n".join(blocks)


def _style(cell, *, bold=False, color=None, fill=None, align=CENTER, border=True):
    cell.font = Font(name=FONT, size=11, bold=bold, color=color)
    cell.alignment = align
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        cell.border = BORDER


def _add_signature(ws, path: Path | None, anchor: str) -> None:
    if not path or not path.exists():
        return
    with PILImage.open(path) as im:
        w, h = im.size
    scale = min(SIGNATURE_MAX_PX[0] / w, SIGNATURE_MAX_PX[1] / h, 1)
    img = XLImage(str(path))
    img.width, img.height = int(w * scale), int(h * scale)
    ws.add_image(img, anchor)


def render_timesheet(
    draft: TimesheetDraft,
    template: TemplateSpec,
    *,
    employee_signature: Path | None = None,
    team_lead_signature: Path | None = None,
) -> bytes:
    if getattr(template, "layout_type", "default") == "adi":
        return _render_adi_timesheet(
            draft,
            template,
            employee_signature=employee_signature,
            team_lead_signature=team_lead_signature,
        )
    return _render_default_timesheet(
        draft,
        template,
        employee_signature=employee_signature,
        team_lead_signature=team_lead_signature,
    )


def _render_default_timesheet(
    draft: TimesheetDraft,
    template: TemplateSpec,
    *,
    employee_signature: Path | None = None,
    team_lead_signature: Path | None = None,
) -> bytes:
    lang = template.language
    wb = Workbook()
    ws = wb.active
    ws.title = i18n.month_name(lang, draft.month)
    ws.sheet_view.showGridLines = True
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    # Baris 1-7: area logo template referensi (kolom B:D)
    company_logo_name = template.company_logo or (
        "pasti_logo.png" if "PASTI" in template.company_name.upper() else None
    )
    if company_logo_name:
        logo_path = ASSETS_DIR / company_logo_name
        if logo_path.exists():
            img0 = XLImage(str(logo_path))
            img0.width, img0.height = 138, 135
            ws.add_image(img0, "B1")

    client_logo_name = template.client_logo or (
        "siloam_logo.png" if "SILOAM" in template.client.upper() else None
    )
    if client_logo_name:
        client_logo_path = ASSETS_DIR / client_logo_name
        if client_logo_path.exists():
            img1 = XLImage(str(client_logo_path))
            img1.width, img1.height = 465, 126
            ws.add_image(img1, "D1")

    ws["A9"] = template.company_name
    ws["A9"].font = Font(name=FONT, size=12, bold=True)
    # Header info (baris 11-12)
    ws.merge_cells("A11:B11")
    ws.merge_cells("A12:B12")
    last_day = max(r.date for r in draft.rows) if draft.rows else None
    month_label = i18n.month_name(lang, draft.month)
    period = f"1 {month_label} {draft.year} - {(last_day.day if last_day else 31)} {month_label} {draft.year}"
    headers = {
        "A11": "Title",
        "C11": "Name",
        "D11": "Company",
        "E11": "Dept",
        "F11": i18n.label(lang, "period"),
    }
    values = {
        "A12": template.title,
        "C12": template.employee_name,
        "D12": template.client,
        "E12": template.department,
        "F12": period,
    }
    for col in "ABCDEF":
        _style(ws[f"{col}11"], bold=True, color="FFFFFFFF", fill=BLACK)
        _style(ws[f"{col}12"])
    for ref, text in headers.items():
        ws[ref] = text
    for ref, text in values.items():
        ws[ref] = text

    # Header tabel (baris 14)
    for i, name in enumerate(
        ["Id", "Date", "Time In", "Time Out", "Working Hour", "Information Activity"], start=1
    ):
        cell = ws.cell(HEADER_ROW_TABLE, i, name)
        _style(cell, bold=True, color="FFFFFFFF", fill=BLACK)

    # Baris data
    r = FIRST_DATA_ROW
    for idx, row in enumerate(draft.rows, start=1):
        for col in range(1, 7):
            _style(ws.cell(r, col), align=LEFT_WRAP if col == 6 else CENTER)
        ws.cell(r, 1, idx)
        ws.cell(r, 2, row.date).number_format = "d mmm yyyy"

        if row.kind == "leave":
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
            ws.cell(r, 3, i18n.label(lang, "leave"))
            _style(ws.cell(r, 3), bold=True, color="FFFFFFFF", fill=RED)
        elif row.kind in ("weekend", "holiday"):
            for col in range(3, 7):
                _style(
                    ws.cell(r, col),
                    bold=True,
                    color="FFFFFFFF",
                    fill=RED,
                    align=LEFT_WRAP if col == 6 else CENTER,
                )
        else:
            if row.time_in and row.time_out:
                ws.cell(r, 3, _t(row.time_in)).number_format = "hh:mm"
                ws.cell(r, 4, _t(row.time_out)).number_format = "hh:mm"
                hours = row.working_hours
                if hours is None:
                    hours = default_working_hours(row.time_in, row.time_out, template.break_minutes)
                # Ditulis sebagai angka (bukan rumus): openpyxl tidak menghitung nilai rumus.
                ws.cell(r, 5, hours).number_format = "0.##"
            text = info_text(row)
            ws.cell(r, 6, text)
            line_count = sum(max(1, math.ceil(len(line) / 95)) for line in text.split("\n")) if text else 1
            ws.row_dimensions[r].height = max(18, 15 * line_count)
        r += 1

    # Footer: tanda tangan
    last_row = r - 1
    label_row = last_row + 2
    ws.cell(label_row, 2, i18n.label(lang, "employee"))
    ws.cell(label_row, 6, i18n.label(lang, "team_lead"))
    for col in (2, 6):
        ws.cell(label_row, col).font = Font(name=FONT, size=11, bold=True)
        ws.cell(label_row, col).alignment = Alignment(horizontal="center" if col == 2 else "left")
    sig_row = label_row + 1
    for rr in range(sig_row, sig_row + 4):
        ws.row_dimensions[rr].height = 20
    _add_signature(ws, employee_signature, f"B{sig_row}")
    _add_signature(ws, team_lead_signature, f"F{sig_row}")
    name_row = sig_row + 4
    ws.cell(name_row, 2, template.employee_name).font = Font(name=FONT, size=11)
    ws.cell(name_row, 6, template.team_lead_name).font = Font(name=FONT, size=11)
    ws.cell(name_row, 6).alignment = Alignment(horizontal="left")

    ws.freeze_panes = f"A{FIRST_DATA_ROW}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_adi_timesheet(
    draft: TimesheetDraft,
    template: TemplateSpec,
    *,
    employee_signature: Path | None = None,
    team_lead_signature: Path | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Absensi Kehadiran-Time Sheet"
    ws.sheet_view.showGridLines = True

    col_widths = {
        "A": 13.88,
        "B": 13.0,
        "C": 11.13,
        "D": 11.88,
        "E": 12.75,
        "F": 13.25,
        "G": 19.0,
        "H": 10.63,
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Logo A1
    logo_file = template.company_logo or "adi_logo.png"
    logo_path = ASSETS_DIR / logo_file
    if logo_path.exists():
        img = XLImage(str(logo_path))
        img.width, img.height = 122, 112
        ws.add_image(img, "A1")

    # Header Alamat F1:H6
    address_lines = [
        (1, "PT. ADI Consulting", 10, True),
        (2, "Gedung PPHUI, 2nd Floor - Suite 210", 9, False),
        (3, "Jl. H.R. Rasuna Said Kav. C 22", 9, False),
        (4, "Jakarta Selatan 12940", 9, False),
        (5, "T: +62 21 278 99868  F: +62 21 2252848", 9, False),
        (6, "www.adiresourcing.com", 9, False),
    ]
    for row_idx, text, font_size, is_bold in address_lines:
        ws.merge_cells(f"F{row_idx}:H{row_idx}")
        cell = ws[f"F{row_idx}"]
        cell.value = text
        cell.font = Font(name="Calibri", size=font_size, bold=is_bold)
        cell.alignment = Alignment(horizontal="right")

    # Baris 9: Banner Judul
    ws.merge_cells("A9:H9")
    ws["A9"] = "Absensi Kehadiran / Time Sheet"
    ws["A9"].font = Font(name="Calibri", size=20, bold=True, color="FFFFFFFF")
    ws["A9"].fill = PatternFill("solid", fgColor="FF7F7F7F")
    ws["A9"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[9].height = 32

    # Baris 10-11: Data Karyawan & Total Hari Kerja
    workdays_count = len([r for r in draft.rows if r.kind == "work"])
    ws["A10"] = "Nama Karyawan / Employee Name: "
    ws["A10"].font = Font(name="Calibri", size=9, bold=True)
    ws["C10"] = template.employee_name
    ws["C10"].font = Font(name="Calibri", size=9, bold=True, underline="single")
    ws["F10"] = "Total Hari Kerja:   "
    ws["F10"].font = Font(name="Calibri", size=9, bold=True)
    ws["G10"] = f"{workdays_count} hari"
    ws["G10"].font = Font(name="Calibri", size=9)

    ws["A11"] = "Divisi / Department                           :"
    ws["A11"].font = Font(name="Calibri", size=9, bold=True)
    ws["C11"] = template.department
    ws["C11"].font = Font(name="Calibri", size=9)
    ws["F11"] = "Total Working Day"
    ws["F11"].font = Font(name="Calibri", size=8, italic=True)

    # Baris 13: Header Tabel
    headers = [
        ("A", "Tanggal / Date\n"),
        ("B", "Waktu Mulai Bekerja / Start Time"),
        ("C", "Waktu Selesai Bekerja / End Time"),
        ("D", "Jumlah Waktu Bekerja / Total Working Hour \n(Hrs) "),
        ("E", "Jam Lembur / Overtime Hour \n(Hrs) "),
        ("F", "Detil Pekerjaan / Job Details"),
        ("H", "Keterangan / Remarks"),
    ]
    ws.row_dimensions[13].height = 36
    ws.merge_cells("F13:G13")
    for col_letter, text in headers:
        ws[f"{col_letter}13"] = text
    for col in "ABCDEFGH":
        c = ws[f"{col}13"]
        c.font = Font(name="Calibri", size=9, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    # Baris Data (mulai baris 14)
    r = 14
    total_hours = 0.0
    for row in draft.rows:
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
        for c in range(1, 9):
            ws.cell(r, c).border = border
            ws.cell(r, c).font = Font(name="Calibri", size=9)

        ws.cell(r, 1, row.date).number_format = "mm/dd/yyyy"
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")

        t_in = row.time_in or template.work_start
        t_out = row.time_out or template.work_end
        hours = (
            row.working_hours
            if row.working_hours is not None
            else default_working_hours(t_in, t_out, template.break_minutes)
        )
        total_hours += hours

        h_in, m_in = t_in.split(":")
        h_out, m_out = t_out.split(":")
        ws.cell(r, 2, time(int(h_in), int(m_in))).number_format = "hh:mm"
        ws.cell(r, 2).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 3, time(int(h_out), int(m_out))).number_format = "hh:mm"
        ws.cell(r, 3).alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(r, 4, hours).number_format = "0.00"
        ws.cell(r, 4).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 5, 0.0).number_format = "0.00"
        ws.cell(r, 5).alignment = Alignment(horizontal="center", vertical="center")

        text = info_text(row)
        ws.cell(r, 6, text)
        ws.cell(r, 6).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        line_count = sum(max(1, math.ceil(len(line) / 45)) for line in text.split("\n")) if text else 1
        ws.row_dimensions[r].height = max(18, 15 * line_count)

        remarks = (
            "CUTI" if row.kind == "leave" else ("LIBUR" if row.kind in ("holiday", "weekend") else "WFO")
        )
        ws.cell(r, 8, remarks).alignment = Alignment(horizontal="center", vertical="center")
        r += 1

    # Baris Totals
    ws.cell(r, 3, "Totals").font = Font(name="Calibri", size=9, bold=True)
    ws.cell(r, 3).alignment = Alignment(horizontal="right", vertical="center")
    ws.cell(r, 4, total_hours).number_format = "0.00"
    ws.cell(r, 4).font = Font(name="Calibri", size=9, bold=True)
    ws.cell(r, 4).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(r, 5, 0.0).number_format = "0.00"
    ws.cell(r, 5).font = Font(name="Calibri", size=9, bold=True)
    ws.cell(r, 5).alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, 9):
        ws.cell(r, c).border = border
    r += 1

    # Baris Tanda Tangan Footer
    ws.cell(r + 1, 1, "Tanda-tangan Karyawan:").font = Font(name="Calibri", size=10)
    ws.cell(r + 2, 1, "Employee Signature").font = Font(name="Calibri", size=9, italic=True)
    ws.cell(r + 1, 5, "Tanda-tangan Kepala Divisi:").font = Font(name="Calibri", size=10, bold=True)
    ws.cell(r + 2, 5, "Supervisor Signature").font = Font(name="Calibri", size=9, italic=True)

    sig_row = r + 3
    for rr in range(sig_row, sig_row + 3):
        ws.row_dimensions[rr].height = 20
    _add_signature(ws, employee_signature, f"A{sig_row}")
    _add_signature(ws, team_lead_signature, f"E{sig_row}")

    name_row = sig_row + 3
    ws.cell(name_row, 1, template.employee_name).font = Font(
        name="Calibri", size=10, bold=True, underline="single"
    )
    ws.cell(name_row, 5, template.team_lead_name).font = Font(
        name="Calibri", size=10, bold=True, underline="single"
    )

    last_d = draft.rows[-1].date if draft.rows else date(draft.year, draft.month, 28)
    date_row = name_row + 1
    ws.cell(date_row, 1, "Tanggal / Date: ").font = Font(name="Calibri", size=9)
    ws.cell(date_row, 2, last_d).number_format = "dd/mm/yyyy"
    ws.cell(date_row, 2).font = Font(name="Calibri", size=9)
    ws.cell(date_row, 5, "Tanggal / Date: ").font = Font(name="Calibri", size=9)
    ws.cell(date_row, 6, last_d).number_format = "dd/mm/yyyy"
    ws.cell(date_row, 6).font = Font(name="Calibri", size=9)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_name(template: TemplateSpec, year: int) -> str:
    safe = template.employee_name.strip().replace(" ", "_")
    return template.file_name_pattern.format(name=safe, year=year)
