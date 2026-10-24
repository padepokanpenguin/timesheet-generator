import io
from datetime import date, time

from openpyxl import load_workbook
from PIL import Image
from timesheet_core.excel_renderer import file_name, info_text, render_timesheet
from timesheet_core.preview import build_draft
from timesheet_core.sample_data import sample_month
from timesheet_core.templates import get_template


def make(template_id="default-id", leave=None, **kw):
    t = get_template(template_id)
    d = build_draft(year=2026, month=1, template=t, data=sample_month(2026, 1), leave_dates=leave or set())
    wb = load_workbook(io.BytesIO(render_timesheet(d, t, **kw)))
    return t, d, wb


def test_layout_matches_reference():
    t, d, wb = make()
    ws = wb["Januari"]
    assert ws["A9"].value == "PT. PASTI MAKMUR ABADI" and ws["A9"].font.b
    assert ws["F11"].value == "Periode"
    assert ws["F12"].value == "1 Januari 2026 - 31 Januari 2026"
    assert [c.value for c in ws[14]][:6] == [
        "Id",
        "Date",
        "Time In",
        "Time Out",
        "Working Hour",
        "Information Activity",
    ]
    assert "A11:B11" in {str(r) for r in ws.merged_cells.ranges}
    assert ws["A14"].fill.fgColor.rgb == "FF000000" and ws["A14"].font.color.rgb == "FFFFFFFF"
    assert ws["A14"].font.name == "Arial"
    assert ws.column_dimensions["F"].width > 90


def test_data_rows_formats_and_colors():
    _, d, wb = make(leave={date(2026, 1, 14)})
    ws = wb["Januari"]
    assert ws.cell(15 + 30, 1).value == 31  # 31 baris data
    assert ws["B15"].number_format == "d mmm yyyy"
    assert ws["C15"].fill.fgColor.rgb == "FFFF0000"  # 1 Jan libur
    # 14 Jan = cuti -> merge C:F, teks CUTI
    row = 15 + 13
    assert ws.cell(row, 3).value == "CUTI"
    assert f"C{row}:F{row}" in {str(r) for r in ws.merged_cells.ranges}
    # hari kerja terisi
    row21 = 15 + 20
    assert ws.cell(row21, 3).value == time(8, 30) and ws.cell(row21, 5).value == 8
    assert ws.cell(row21, 6).value.startswith("[Stabilization Kairos] BPAS-5437:")
    assert ws.cell(row21, 6).alignment.wrap_text


def test_info_text_format():
    _, d, _ = make()
    row = next(r for r in d.rows if r.date == date(2026, 1, 14))
    text = info_text(row)
    assert "[Kairos Pas] KP2-68: FE - Setup design system package" in text
    assert " - setup package and build config" in text
    assert "[Stabilization Kairos] BPAS-5115" in text
    # Verifikasi pemisahan baris baru (newline \n) untuk task ke-2
    lines = text.split("\n")
    assert len(lines) == 2
    assert lines[0].startswith("[Stabilization Kairos] BPAS-5115")
    assert lines[1].startswith("[Kairos Pas] KP2-68")

def test_english_template_and_footer(tmp_path):
    sig = tmp_path / "sig.png"
    Image.new("RGBA", (300, 120), (0, 0, 0, 0)).save(sig)
    t, d, wb = make("default-en", employee_signature=sig, team_lead_signature=sig)
    ws = wb["January"]
    assert ws["F11"].value == "Period"
    assert ws["F12"].value == "1 January 2026 - 31 January 2026"
    footer = [(c.coordinate, c.value) for row in ws.iter_rows(min_row=46) for c in row if c.value]
    assert ("B47", "Employee") in footer and ("F47", "Team Lead") in footer
    assert ("B52", "Aria Nur Jamal") in footer and ("F52", "Kevin Saputra") in footer
    assert len(ws._images) == 4


def test_file_name():
    assert file_name(get_template("default-id"), 2026) == "TimeSheet_Aria_Nur_Jamal_2026.xlsx"
